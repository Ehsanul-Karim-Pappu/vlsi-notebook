#!/usr/bin/env python3
"""Render lecture 1 and build its slide deck: English slides, English and Bangla speaker notes.

    python build.py                  draft render (480p, 15 fps) of every chapter
    python build.py -q p             for presenting: 1080p, 30 fps
    python build.py --only Ch05      render only the matching chapters, then rebuild the deck
    python build.py --deck-only      rebuild the deck from what's already rendered
    python build.py --core           also build the 25-minute core deck, build/deck/core.html
    python build.py --pptx           also export PowerPoint (labs as screenshots; with --core, both)
    python build.py --movie          also cut a continuous movie, with subtitles (--core, --lang bn)
    python build.py --cdn            load reveal.js from the internet instead of bundling it

Quality: l 480p15 (draft), m 720p30, p 1080p30 (presenting), h 1080p60, k 4K60.

The deck lands in build/deck/. Open index.html in a browser: → and ← move between slides, ↓
opens a derivation where there is one, S opens the speaker view, F goes full screen. core.html,
next to it, is the 25-minute core path (core.py), using the same videos. Run it from anywhere;
paths are relative to this file.
"""

import argparse
import base64
import json
import os
import re
import tempfile
import textwrap
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
FONTS = HERE.parents[1] / "kit" / "fonts"
BUILD = HERE / "build"

TITLE = "Lecture 1 · The Switch That Wouldn't Turn Off"
AUTHOR = "Khandaker Ehsanul Karim"  # shown small on every slide: "Prepared by ..."
# Between a note's English and its Bangla, as kit.style.say() writes it.
BANGLA = "— বাংলা —"

# The deck, in order: ("scene", file, class) for a Manim chapter, ("lab", file, title, (English
# notes, Bangla notes)) for a live lab.
SEQUENCE = [
    ("scene", "chapters/ch00_cold_open.py", "Ch00ColdOpen"),
    ("scene", "chapters/ch01_before_the_switch.py", "Ch01BeforeTheSwitch"),
    ("scene", "chapters/ch02_accidental_transistor.py", "Ch02AccidentalTransistor"),
    ("scene", "chapters/ch03_glass.py", "Ch03Glass"),
    ("scene", "chapters/ch04_free_lunch.py", "Ch04FreeLunch"),
    (
        "lab",
        "lab3_dennard.html",
        "Live lab 3: Dennard's dial",
        (
            "Live lab. Press Shrink 4 generations: sixteen times the transistors in the patch, each "
            "about four times as fast, and the power density stays at one. That's Dennard's "
            "constant-field scaling, the free lunch. Now press Voltage stuck, clock free: the same "
            "shrink, but the voltage stays put, and the power density shoots up. Then Voltage stuck, "
            "clock held: it still climbs, more slowly. That's the trouble the next chapter is about. "
            "It's an ideal model: no leakage, no wires.",
            "এইটা live lab। Shrink 4 generations চাপেন: patch-এ ষোল গুণ transistor, প্রত্যেকটা প্রায় চার গুণ "
            "দ্রুত, আর power density এক-এই থাকে। এইটা Dennard-এর constant-field scaling, free lunch। এবার "
            "Voltage stuck, clock free চাপেন: একই shrink, কিন্তু voltage একই জায়গায়, আর power density লাফ দিয়া "
            "বাড়ে। তারপর Voltage stuck, clock held: তবুও বাড়ে, একটু ধীরে। পরের chapter এই ঝামেলা নিয়াই। এইটা "
            "ideal model: leakage নাই, wire নাই।",
        ),
    ),
    ("scene", "chapters/ch05_boltzmann.py", "Ch05Boltzmann"),
    (
        "lab",
        "lab1_boltzmann.html",
        "Live lab 1: Boltzmann's fence",
        (
            "Live lab. Drag the gate voltage slowly up from zero and watch electrons start "
            "to cross the hill; the dot on the right climbs one decade for every 60 mV. Then press "
            "77 K: in this ideal model the swing drops to about 15 mV per decade (real cold "
            "transistors don't quite get there). Then raise the grip, m, to 1.4: the slope gets "
            "lazier. The moving electrons are an illustration; the plot is the model. That's the "
            "whole chapter on one screen.",
            "এইটা live lab। Gate voltage শূন্য থেকে আস্তে আস্তে বাড়ান, দেখবেন electron-গুলা hill পার হওয়া "
            "শুরু করছে, আর ডানের dot-টা প্রতি 60 mV-তে এক decade করে উঠতেছে। এরপর 77 K চাপেন: এই ideal "
            "model-এ swing নেমে আসে প্রায় 15 mV/decade-এ (আসল ঠান্ডা transistor পুরাটা পৌঁছায় না)। তারপর "
            "grip, মানে m, বাড়ায়ে 1.4 করেন: slope-টা ঢিলা হয়ে যায়। চলন্ত electron-গুলা শুধু ছবি; plot-টা "
            "model। পুরা chapter-টা এক screen-এ।",
        ),
    ),
    ("scene", "chapters/ch06_losing_grip.py", "Ch06LosingGrip"),
    (
        "lab",
        "lab2_barrier.html",
        "Live lab 2: Who controls the barrier?",
        (
            "Live lab. It opens on a planar transistor with a 30 nm gate: about four and a half "
            "lambda, in the red. Raise the drain voltage and watch the hill sink and the leakage "
            "climb. Now press FinFET: the same gate length, but lambda drops to 2.4 nm and the gate "
            "is back in control. Try the nanosheet and the 2D layer. Then press Shrink the gate and "
            "watch where each one gives up.",
            "এইটা live lab। শুরু হয় 30 nm gate-এর একটা planar transistor দিয়া: প্রায় সাড়ে চার lambda, লাল "
            "অংশে। Drain voltage বাড়ান, দেখেন hill নেমে যায় আর leakage বাড়ে। এবার FinFET চাপেন: gate length "
            "একই, কিন্তু lambda নেমে 2.4 nm, আর gate আবার control-এ। Nanosheet আর 2D layer try করেন। তারপর "
            "Shrink the gate চাপেন, দেখেন কোনটা কোথায় হাল ছাড়ে।",
        ),
    ),
    ("scene", "chapters/ch07_finfet.py", "Ch07FinFET"),
    ("scene", "chapters/ch08_nanosheet.py", "Ch08Nanosheet"),
    ("scene", "chapters/ch09_forksheet.py", "Ch09Forksheet"),
    ("scene", "chapters/ch10_cfet.py", "Ch10CFET"),
    (
        "lab",
        "lab4_cell.html",
        "Live lab 4: Build a cell",
        (
            "Live lab. It opens on a FinFET inverter, two fins per transistor: it needs seven "
            "tracks of 24 nanometres. Switch to the nanosheet and press Fewest tracks: one 22 "
            "nanometre sheet, six tracks. The forksheet's wall: five. The CFET, n stacked on p: four. "
            "Widen the sheets and watch the cell grow; the right side shows the effective width you "
            "get for it. These are schematic models with FET Lab's spacings, not a foundry's rules.",
            "এইটা live lab। শুরু হয় একটা FinFET inverter দিয়া, প্রতি transistor-এ দুইটা fin: 24 nanometre-এর "
            "সাতটা track লাগে। Nanosheet-এ যান আর Fewest tracks চাপেন: 22 nanometre-এর একটা sheet, ছয়টা track। "
            "Forksheet-এর wall দিয়া: পাঁচটা। CFET, p-এর উপরে n: চারটা। Sheet চওড়া করেন, দেখেন cell বড় হয়; ডান "
            "দিকে দেখায় এর বদলে কতটা effective width পাইতেছেন। এগুলা FET Lab-এর spacing দিয়া schematic model, "
            "কোনো foundry-র rule না।",
        ),
    ),
    (
        "lab",
        "lab5_models.html",
        "Live lab 5: Hold it yourself",
        (
            "Live lab. These are FET Lab's own 3D models. Drag to turn one, scroll to zoom. Press Cut "
            "through the gate and look at the cut face. Then open the monolithic CFET, cut it through "
            "the gate, and switch layers off on the right to see the sheets inside. Hover over any "
            "piece to see what it is.",
            "এইটা live lab। এগুলা FET Lab-এর নিজের 3D model। টেনে ঘুরান, scroll করে zoom করেন। Cut through the "
            "gate চাপেন, কাটা মুখটা দেখেন। তারপর monolithic CFET খোলেন, gate বরাবর কাটেন, আর ডানের layer-গুলা বন্ধ "
            "করে ভিতরের sheet দেখেন। কোনো অংশের উপর pointer রাখলে নাম দেখাবে।",
        ),
    ),
    ("scene", "chapters/ch11_beyond.py", "Ch11Beyond"),
    ("scene", "chapters/ch12_outro.py", "Ch12Outro"),
]

# Manim settings per quality. Use long options: manim-slides reads the "h" in "-qh" as its own
# -h (help) and exits without rendering.
QUALITY = {
    "l": (["--quality", "l"], "480p, 15 fps"),
    "m": (["--quality", "m"], "720p, 30 fps"),
    "p": (["--resolution", "1920,1080", "--frame_rate", "30"], "1080p, 30 fps"),
    "h": (["--quality", "h"], "1080p, 60 fps"),
    "k": (["--quality", "k"], "4K, 60 fps"),
}


def manim_slides(*args, **kw):
    return subprocess.run([sys.executable, "-m", "manim_slides", *args], cwd=HERE, **kw)


def render(scene_file, scene, quality):
    log = BUILD / f"{scene}.log"
    start = time.time()
    with open(log, "w") as out:
        proc = manim_slides(
            # Each scene gets its own media folder: renders run in parallel, and a shared text
            # cache lets one process read a file another is still writing.
            "render", "--media_dir", str(BUILD / "media" / scene), *QUALITY[quality][0], scene_file, scene,
            stdout=out, stderr=subprocess.STDOUT,
        )
    took = time.time() - start
    out_json = HERE / "slides" / f"{scene}.json"
    fresh = out_json.exists() and out_json.stat().st_mtime >= start
    ok = proc.returncode == 0 and fresh
    status = "ok" if ok else f"FAILED (see {log.relative_to(HERE)})"
    print(f"  {scene:<28} {took:6.0f} s  {status}", flush=True)
    return ok


def lab_section(name, title, notes):
    en, bn = notes
    return (
        '<section class="lab" data-background-color="#0E1116">'
        f'<iframe data-src="labs/{name}" title="{title}" allow="fullscreen"></iframe>'
        f'<aside class="notes">{en}\n\n{BANGLA}\n\n{bn}</aside>'
        "</section>"
    )


def models_js():
    """FET Lab's 3D models (models/*.glb), base64 in one script, so Live Lab 5 can load them from
    a file:// deck, where a browser won't fetch() a local file. Rewritten only when it changes."""
    lines = ["// FET Lab's 3D models (models/*.glb), base64. Written by build.py; don't edit.",
             "window.LAB5_MODELS = {"]
    for glb in sorted((HERE / "models").glob("*.glb")):
        lines.append(f'  {glb.stem}: "{base64.b64encode(glb.read_bytes()).decode()}",')
    lines.append("};")
    text = "\n".join(lines) + "\n"
    out = HERE / "labs" / "lab5-models.js"
    if not out.exists() or out.read_text() != text:
        out.write_text(text)


def thumbnails(deck):
    """The last frame of every slide's video, small, for the speaker view's "Next" preview."""
    out = deck / "thumbs"
    out.mkdir()
    for video in sorted((deck / "index_assets").glob("*.mp4")):
        jpg = out / (video.stem + ".jpg")
        for seek in (["-sseof", "-0.1"], []):  # a very short clip has no 0.1 s to seek back
            subprocess.run(["ffmpeg", "-v", "error", "-y", *seek, "-i", str(video), "-frames:v", "1",
                            "-vf", "scale=640:-2", "-q:v", "4", str(jpg)], check=False)
            if jpg.exists():
                break


def scene_names():
    return [s[2] for s in SEQUENCE if s[0] == "scene"]


def core_folder():
    """slides/ cut down to the core path (core.py), with its shorter notes: a folder for
    manim-slides convert --folder. Every scene keeps at least one slide, so the videos' file names
    (prefixed with the scene's number) match the full deck's."""
    from core import CORE

    out = BUILD / "core-slides"
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    for scene in scene_names():
        config = json.loads((HERE / "slides" / f"{scene}.json").read_text())
        kept = []
        for i, anchor, en, bn in CORE.get(scene, []):
            slide = config["slides"][i]
            full = " ".join(slide["notes"].split(BANGLA)[0].split())
            if not full.startswith(anchor):
                sys.exit(f"core.py: {scene} slide {i} no longer starts with {anchor!r}; it starts {full[:40]!r}")
            # manim-slides reads the video paths relative to the slides folder's parent: make them absolute.
            files = {k: str(HERE / slide[k]) for k in ("file", "rev_file") if slide.get(k)}
            kept.append(dict(slide, **files, notes=f"{en}\n\n{BANGLA}\n\n{bn}", direction="horizontal"))
        if not kept:
            sys.exit(f"core.py keeps no slide of {scene}; keep at least one")
        config["slides"] = kept
        (out / f"{scene}.json").write_text(json.dumps(config, ensure_ascii=False))
    return out


def lab_items(core):
    """The labs, in deck order, as (scene index they follow, file, title, (en, bn) notes)."""
    from core import LABS

    items, n_scene = [], -1
    for item in SEQUENCE:
        if item[0] == "scene":
            n_scene += 1
        else:
            name, title, notes = item[1:]
            items.append((n_scene, name, title, LABS[name] if core else notes))
    return items


def convert_html(dest, folder, title, cdn, core):
    """One reveal.js page from the rendered slides, with the labs put after their chapters."""
    manim_slides(
        "convert", *scene_names(), str(dest),
        "--folder", str(folder),
        "--to", "html",
        "--use-template", str(HERE / "template" / "deck.html"),
        f"-ctitle={title}",
        # The core page shares the full deck's videos.
        "-cassets_dir=index_assets",
        "-ccontrols=true",
        "-cprogress=true",
        "-cslide_number=c/t",
        "-chash=true",
        "-cview_distance=2",
        # Bundle reveal.js into the deck, so it presents without an internet connection.
        *([] if cdn else ["--offline"]),
        check=True,
    )
    # Put each lab after the chapter it follows. The template leaves a marker after every
    # scene: <!-- after-scene:N --> with N counting scenes from 0.
    html = dest.read_text().replace("__AUTHOR__", AUTHOR)
    labs_after = {}
    for n, name, lab_title, notes in lab_items(core):
        labs_after.setdefault(n, []).append(lab_section(name, lab_title, notes))
    for n, sections in labs_after.items():
        marker = f"<!-- after-scene:{n} -->"
        if marker not in html:
            sys.exit(f"Template marker {marker} not found")
        html = html.replace(marker, marker + "\n" + "\n".join(sections))
    dest.write_text(html)


def build_deck(pptx=False, cdn=False, core=False):
    deck = BUILD / "deck"
    scenes = scene_names()
    missing = [s for s in scenes if not (HERE / "slides" / f"{s}.json").exists()]
    if missing:
        sys.exit(f"Not rendered yet: {', '.join(missing)}. Run build.py without --deck-only first.")
    if deck.exists():
        shutil.rmtree(deck)
    deck.mkdir(parents=True)
    convert_html(deck / "index.html", HERE / "slides", TITLE, cdn, core=False)
    if core:
        convert_html(deck / "core.html", core_folder(), TITLE + " (core, 25 min)", cdn, core=True)
    models_js()
    shutil.copytree(HERE / "labs", deck / "labs")
    shutil.copy(HERE / "template" / "speaker.html", deck / "speaker.html")
    thumbnails(deck)
    # The fonts, so the labs and the speaker view (whose notes are also in Bangla) look right
    # offline. Both look in labs/fonts/.
    (deck / "labs" / "fonts").mkdir()
    for f in [*FONTS.glob("Inter-*.ttf"), *FONTS.glob("NotoSansBengali-*.ttf"), *FONTS.glob("OFL*.txt")]:
        shutil.copy(f, deck / "labs" / "fonts" / f.name)
    print(f"Deck: {(deck / 'index.html').relative_to(HERE)}" + (f", {(deck / 'core.html').relative_to(HERE)}" if core else ""))
    if pptx:
        export_pptx(BUILD / "lecture1.pptx", HERE / "slides", core=False)
        if core:
            export_pptx(BUILD / "lecture1-core.pptx", BUILD / "core-slides", core=True)


# --- PowerPoint ------------------------------------------------------------------------------
LAB_SHOTS = HERE / "images" / "labs"  # a screenshot of each lab, for PowerPoint and the movie


def export_pptx(out, folder, core):
    """PowerPoint, with the animations as videos and the notes in both languages. A lab can't
    run in PowerPoint, so its slide is a screenshot that says where the live one is."""
    from pptx import Presentation
    from pptx.util import Pt

    manim_slides("convert", *scene_names(), str(out), "--folder", str(folder), "--to", "pptx", check=True)
    prs = Presentation(str(out))
    counts = [len(json.loads((folder / f"{s}.json").read_text())["slides"]) for s in scene_names()]
    ids = prs.slides._sldIdLst
    # From the last lab back, so the earlier positions don't move.
    for n, name, title, (en, bn) in reversed(lab_items(core)):
        slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1])
        shot = LAB_SHOTS / (Path(name).stem + ".jpg")
        if shot.exists():
            slide.shapes.add_picture(str(shot), 0, 0, prs.slide_width, prs.slide_height)
        box = slide.shapes.add_textbox(Pt(18), prs.slide_height - Pt(40), prs.slide_width - Pt(36), Pt(30))
        box.text_frame.text = f"{title}: a screenshot. Open the HTML deck (build/deck/index.html) to use it live."
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        slide.notes_slide.notes_text_frame.text = f"{en}\n\n{BANGLA}\n\n{bn}"
        ids.insert(sum(counts[: n + 1]), ids[-1])
    prs.save(str(out))
    print(f"PowerPoint: {out.relative_to(HERE)} ({len(prs.slides)} slides)")


# --- the continuous cut --------------------------------------------------------------------------
NARRATION = HERE / "narration"
WPM = 150  # speaking pace for slides with no recording yet
LEAD, TAIL = 0.25, 0.6  # seconds of quiet before and after each recording


def probe_seconds(path):
    out = subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0", str(path)],
                         capture_output=True, text=True, check=True).stdout
    return float(out.strip())


def untagged(text):
    """A narration line without its audio tags ([curious], [short pause]), as the subtitles show it."""
    return re.sub(r"\s+([,.;:?!])", r"\1", " ".join(re.sub(r"\[[^\]]*\]", " ", text).split()))


def recording(folder, stem):
    for ext in (".wav", ".m4a", ".mp3", ".flac", ".ogg"):
        if (folder / (stem + ext)).exists():
            return folder / (stem + ext)
    return None


def srt_time(t):
    ms = int(round(t * 1000))
    return f"{ms // 3600000:02d}:{ms // 60000 % 60:02d}:{ms // 1000 % 60:02d},{ms % 1000:03d}"


def movie(core=False, lang="en"):
    """A continuous movie of the deck, for a recorded version of the talk.

    Each slide plays its animation and then holds (or keeps looping) for as long as its narration
    lasts. Narration is yours: record one file per slide into narration/<lang>/ (narration/<lang>-core/
    for the core path), named as in the script this writes next to the movie, and build again.
    A slide with no recording yet is held for its notes' length at 150 words a minute, silent, so the
    timing is right to record against. A lab shows its screenshot; put a screen recording of you
    working it at narration/<lang>/<lab name>.mp4 (with its own sound) and it is used instead.
    Writes build/lecture1[-core]-<lang>.mp4, a matching .srt subtitle file, and the script."""
    folder = core_folder() if core else HERE / "slides"
    tag = f"{'-core' if core else ''}-{lang}"
    rec = NARRATION / (lang + ("-core" if core else ""))
    first = json.loads((folder / f"{scene_names()[0]}.json").read_text())
    width, height = first["resolution"]
    labs_after = {}
    for n, name, title, notes in lab_items(core):
        labs_after.setdefault(n, []).append((name, title, notes))
    from core import CORE

    items = []  # (stem, kind, source, notes text, loop)
    for n, scene in enumerate(scene_names()):
        config = json.loads((folder / f"{scene}.json").read_text())
        # Recordings are named by the slide's number in the full scene, core or not.
        numbers = [row[0] for row in CORE[scene]] if core else range(len(config["slides"]))
        for i, slide in zip(numbers, config["slides"]):
            en, _, bn = slide["notes"].partition(BANGLA)
            items.append((f"{scene}-{i:02d}", "video", HERE / slide["file"], (en.strip(), bn.strip()), slide["loop"]))
        for name, title, notes in labs_after.get(n, []):
            items.append((Path(name).stem, "lab", LAB_SHOTS / (Path(name).stem + ".jpg"), notes, False))

    # If the recordings come with a manifest (file,seconds,text), the subtitles follow what was said,
    # without the audio tags that directed it.
    said = {}
    if (rec / "manifest.csv").exists():
        import csv

        with open(rec / "manifest.csv", newline="", encoding="utf-8") as f:
            said = {Path(row["file"]).stem: untagged(row["text"]) for row in csv.DictReader(f) if row.get("text")}
    script = [f"# Narration script: {TITLE}{' (core)' if core else ''}, {lang}", "",
              f"Record each block into {rec.relative_to(HERE)}/<file>.wav (or .m4a, .mp3), then run build.py --movie again.", ""]
    srt, t, segments = [], 0.0, []
    work = Path(tempfile.mkdtemp(prefix="cut-", dir=BUILD))
    for k, (stem, kind, src, (en, bn), loop) in enumerate(items):
        text = en if lang == "en" else bn
        script += [f"## {stem}", "", text, ""]
        text = said.get(stem, text)
        screen = rec / (stem + ".mp4") if kind == "lab" else None
        audio = recording(rec, stem)
        clip = probe_seconds(src) if kind == "video" else 0.0
        if screen and screen.exists():
            dur = probe_seconds(screen)
        else:
            spoken = probe_seconds(audio) + LEAD + TAIL if audio else len(en.split()) * 60 / WPM + 1.0
            dur = max(clip, spoken, 2.0)
        seg = work / f"{k:04d}.mp4"
        # The author line, as on every slide of the deck.
        font = str(FONTS / "Inter-Regular.ttf").replace("\\", "/").replace(":", "\\:")  # ffmpeg filter escaping
        credit = (f"drawtext=fontfile='{font}':text='Prepared by {AUTHOR}':"
                  f"fontcolor=0x8A8F98@0.75:fontsize={max(10, height // 55)}:x={height // 45}:y=h-th-{height // 70}")
        vf = (f"scale={width}:{height}:force_original_aspect_ratio=decrease,pad={width}:{height}:(ow-iw)/2:(oh-ih)/2,"
              f"{credit},fps=30,format=yuv420p")
        if screen and screen.exists():
            inputs, vmap, amap = ["-i", str(screen)], f"[0:v]{vf}[v]", "[0:a]aresample=48000,apad[a]"
        else:
            if kind == "lab":
                inputs = ["-loop", "1", "-framerate", "30", "-i", str(src)]
            elif loop:
                inputs = ["-stream_loop", "-1", "-i", str(src)]
            else:
                inputs = ["-i", str(src)]
                vf = f"tpad=stop_mode=clone:stop_duration={max(0.0, dur - clip):.3f}," + vf
            vmap = f"[0:v]{vf}[v]"
            if audio:
                inputs += ["-i", str(audio)]
                # Every recording at the same loudness, after a short breath.
                amap = (f"[1:a]loudnorm=I=-16:TP=-1.5:LRA=11,aresample=48000,"
                        f"adelay={int(LEAD * 1000)}:all=1,apad[a]")
            else:
                inputs += ["-f", "lavfi", "-i", "anullsrc=r=48000:cl=stereo"]
                amap = "[1:a]anull[a]"
        subprocess.run(["ffmpeg", "-v", "error", "-y", *inputs, "-filter_complex", f"{vmap};{amap}",
                        "-map", "[v]", "-map", "[a]", "-t", f"{dur:.3f}", "-c:v", "libx264", "-preset", "veryfast",
                        "-crf", "20", "-c:a", "aac", "-b:a", "160k", "-ac", "2", str(seg)], check=True)
        segments.append(seg)
        # Subtitles: the notes, a sentence or two per cue, timed in proportion to their length.
        parts = [p for p in re.split(r"(?<=[.?!।:])\s+", " ".join(text.split())) if p] or [""]
        total = sum(len(p) for p in parts) or 1
        start = t
        for part in parts:
            end = start + dur * len(part) / total
            srt += [str(len(srt) // 4 + 1), f"{srt_time(start)} --> {srt_time(end)}", "\n".join(textwrap.wrap(part, 60)), ""]
            start = end
        t += probe_seconds(seg)
    listing = work / "list.txt"
    listing.write_text("".join(f"file '{seg}'\n" for seg in segments))
    out = BUILD / f"lecture1{tag}.mp4"
    subprocess.run(["ffmpeg", "-v", "error", "-y", "-f", "concat", "-safe", "0", "-i", str(listing), "-c", "copy",
                    "-movflags", "+faststart", str(out)], check=True)
    out.with_suffix(".srt").write_text("\n".join(srt), encoding="utf-8")
    (BUILD / f"narration{tag}.md").write_text("\n".join(script), encoding="utf-8")
    shutil.rmtree(work)
    recorded = sum(1 for stem, *_ in items if recording(rec, stem))
    print(f"Movie: {out.relative_to(HERE)}, {t / 60:.1f} min, {recorded} of {len(items)} slides narrated; "
          f"subtitles {out.with_suffix('.srt').relative_to(HERE)}; script build/narration{tag}.md")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-q", "--quality", choices=QUALITY, default="l")
    ap.add_argument("--only", nargs="*", default=None, help="render only scenes whose class name contains one of these")
    ap.add_argument("--deck-only", action="store_true")
    ap.add_argument("--pptx", action="store_true")
    ap.add_argument("--core", action="store_true", help="also build the 25-minute core path (core.py)")
    ap.add_argument("--movie", action="store_true", help="also cut a continuous movie with subtitles")
    ap.add_argument("--lang", choices=("en", "bn"), default="en", help="the movie's narration and subtitles")
    ap.add_argument("--cdn", action="store_true", help="load reveal.js from the internet instead of bundling it")
    ap.add_argument("-j", "--jobs", type=int, default=os.cpu_count() or 2)
    args = ap.parse_args()

    BUILD.mkdir(exist_ok=True)
    if not args.deck_only:
        todo = [(s[1], s[2]) for s in SEQUENCE if s[0] == "scene"]
        if args.only:
            todo = [t for t in todo if any(k.lower() in t[1].lower() for k in args.only)]
        print(f"Rendering {len(todo)} scene(s) at {QUALITY[args.quality][1]}, {args.jobs} at a time:")
        with ThreadPoolExecutor(args.jobs) as pool:
            ok = list(pool.map(lambda t: render(t[0], t[1], args.quality), todo))
        if not all(ok):
            sys.exit(1)
    build_deck(pptx=args.pptx, cdn=args.cdn, core=args.core)
    if args.movie:
        movie(core=False, lang=args.lang)
        if args.core:
            movie(core=True, lang=args.lang)


if __name__ == "__main__":
    main()
