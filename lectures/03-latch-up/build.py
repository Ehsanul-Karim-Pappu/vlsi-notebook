#!/usr/bin/env python3
"""Render lecture 3 and build its slide deck: English slides, English and Bangla speaker notes.

    python build.py                  draft render (480p, 15 fps) of every chapter
    python build.py -q p             for presenting: 1080p, 30 fps
    python build.py --only X04       render only the matching chapters, then rebuild the deck
    python build.py --deck-only      rebuild the deck from what's already rendered
    python build.py --pptx           also export PowerPoint (the labs as screenshots)
    python build.py --cdn            load reveal.js from the internet instead of bundling it

Quality: l 480p15 (draft), m 720p30, p 1080p30 (presenting), h 1080p60, k 4K60.

The deck lands in build/deck/. Open index.html in a browser: → and ← move between slides, S
opens the speaker view, F goes full screen. Run it from anywhere; paths are relative to this file.
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
FONTS = HERE.parents[1] / "kit" / "fonts"
BUILD = HERE / "build"
DECK = "deck"

TITLE = "Lecture 3 · Latch-up, in Layout"
AUTHOR = "Khandaker Ehsanul Karim"  # shown small on every slide: "Prepared by ..."
# Between a note's English and its Bangla, as kit.style.say() writes it.
BANGLA = "— বাংলা —"

FIRE_LAB = (
    "lab",
    "lab1_fire.html",
    "Live lab: Fire the thyristor",
    (
        "Live lab. The cross-section on the left is the inverter from chapter one, with its two parasitic "
        "bipolars and the resistances to the taps; the plot on the right is the thyristor's I–V curve. "
        "Press Inject: a current pulse goes into the substrate. With the taps close, the drop across R_sub "
        "stays below 0.65 volts and nothing happens. Now drag the tap distance out and inject again: the NPN "
        "wakes, the loop fires, and at 3.3 volts it stays latched until you press Power cycle. Try the 0.8 "
        "volt core supply: it fires and lets go, because the supply is below the holding voltage. Finally, "
        "at 1.8 volts, widen the N+ to well spacing past about 10 microns: the holding voltage climbs past "
        "the supply. Widen it further, past about 25, and the loop gain itself falls below one.",
        "এইটা live lab। বামের cross-section-টা chapter one-এর inverter, তার দুইটা parasitic bipolar আর tap পর্যন্ত "
        "resistance সহ; ডানের plot-টা thyristor-এর I–V curve। Inject চাপেন: একটা current pulse substrate-এ ঢোকে। Tap কাছে "
        "থাকলে R_sub-এর উপর drop 0.65 volt-এর নিচে থাকে, কিছুই হয় না। এবার tap-এর দূরত্ব বাড়ান আর আবার inject করেন: NPN "
        "জেগে ওঠে, loop চালু হয়, আর 3.3 volt-এ Power cycle না চাপা পর্যন্ত latched থাকে। 0.8 volt core supply চেষ্টা করেন: "
        "চালু হয় আর ছেড়ে দেয়, কারণ supply holding voltage-এর নিচে। শেষে, 1.8 volt-এ, N+ থেকে well-এর spacing প্রায় 10 micron-এর "
        "বেশি বাড়ান: holding voltage supply-রে ছাড়িয়ে যায়। আরও বাড়ান, প্রায় 25-এর বেশি, আর loop gain নিজেই একের নিচে নেমে যায়।",
    ),
)

TAP_LAB = (
    "lab",
    "lab2_taps.html",
    "Live lab: Tap the floor",
    (
        "Live lab. A block of transistors from above: N+ active in the substrate, P+ active in two N-wells. "
        "The rule is the one from chapter four: every active must be within the tap distance of a tap of its "
        "own kind. Click a free site to place a tap: a P+ tap in the substrate, an N+ tap in a well. The "
        "shaded squares are each tap's reach, sized exactly the way a DRC sizes them, and every active "
        "outside them gets a red marker. Start from No taps, try Row ends, then fix the block with as few "
        "taps as you can. Turn on Wrong well to see a tap that looks right and counts for nothing.",
        "এইটা live lab। উপর থেকে একটা transistor-এর block: substrate-এ N+ active, দুইটা N-well-এ P+ active। Rule-টা chapter "
        "four-এর: প্রত্যেকটা active তার নিজের ধরনের একটা tap থেকে tap distance-এর মধ্যে থাকতে হবে। একটা খালি জায়গায় click করে "
        "tap বসান: substrate-এ P+ tap, well-এ N+ tap। ছায়া দেওয়া square-গুলা প্রত্যেকটা tap-এর নাগাল, ঠিক যেভাবে একটা DRC "
        "size করে, আর এর বাইরের প্রত্যেকটা active একটা লাল marker পায়। No taps থেকে শুরু করেন, Row ends চেষ্টা করেন, তারপর "
        "যত কম tap দিয়া পারেন block-টা ঠিক করেন। Wrong well চালু করে দেখেন একটা tap যেটা দেখতে ঠিক কিন্তু কোনো কাজে আসে না।",
    ),
)

# The deck, in order: ("scene", file, class) for a Manim chapter, ("lab", file, title, (English
# notes, Bangla notes)) for a live lab.
SEQUENCE = [
    ("scene", "chapters/x00_cold_open.py", "X00ColdOpen"),
    ("scene", "chapters/x01_thyristor.py", "X01Thyristor"),
    ("scene", "chapters/x02_firing.py", "X02Firing"),
    FIRE_LAB,
    ("scene", "chapters/x03_knobs.py", "X03Knobs"),
    ("scene", "chapters/x04_taps.py", "X04Taps"),
    TAP_LAB,
    ("scene", "chapters/x05_guard_rings.py", "X05GuardRings"),
    ("scene", "chapters/x06_injectors.py", "X06Injectors"),
    ("scene", "chapters/x07_wells.py", "X07Wells"),
    ("scene", "chapters/x08_review.py", "X08Review"),
]

# Manim settings per quality. Use long options: manim-slides reads the "h" in "-qh" as its own
# -h (help) and exits without rendering.
QUALITY = {
    "l": (["--quality", "l"], "480p, 15 fps"),
    "m": (["--quality", "m"], "720p, 30 fps"),
    "p": (["--resolution", "1920,1080", "--frame_rate", "30"], "1080p, 30 fps"),
    "h": (["--quality", "h"], "1080p, 60 fps"),
    "k": (["--quality", "k"], "4K, 60 fps"),
}


def manim_slides(*args, **kw):
    return subprocess.run([sys.executable, "-m", "manim_slides", *args], cwd=HERE, **kw)


def render(scene_file, scene, quality):
    log = BUILD / f"{scene}.log"
    start = time.time()
    with open(log, "w") as out:
        proc = manim_slides(
            # Each scene gets its own media folder: renders run in parallel, and a shared text
            # cache lets one process read a file another is still writing.
            "render", "--media_dir", str(BUILD / "media" / scene), *QUALITY[quality][0], scene_file, scene,
            stdout=out, stderr=subprocess.STDOUT,
        )
    took = time.time() - start
    out_json = HERE / "slides" / f"{scene}.json"
    fresh = out_json.exists() and out_json.stat().st_mtime >= start
    ok = proc.returncode == 0 and fresh
    status = "ok" if ok else f"FAILED (see {log.relative_to(HERE)})"
    print(f"  {scene:<28} {took:6.0f} s  {status}", flush=True)
    return ok


def lab_section(name, title, notes):
    en, bn = notes
    return (
        '<section class="lab" data-background-color="#0E1116">'
        f'<iframe data-src="labs/{name}" title="{title}" allow="fullscreen"></iframe>'
        f'<aside class="notes">{en}\n\n{BANGLA}\n\n{bn}</aside>'
        "</section>"
    )


def thumbnails(deck):
    """The last frame of every slide's video, small, for the speaker view's "Next" preview."""
    out = deck / "thumbs"
    out.mkdir()
    for video in sorted((deck / "index_assets").glob("*.mp4")):
        jpg = out / (video.stem + ".jpg")
        for seek in (["-sseof", "-0.1"], []):  # a very short clip has no 0.1 s to seek back
            subprocess.run(["ffmpeg", "-v", "error", "-y", *seek, "-i", str(video), "-frames:v", "1",
                            "-vf", "scale=640:-2", "-q:v", "4", str(jpg)], check=False)
            if jpg.exists():
                break


def scene_names():
    return [s[2] for s in SEQUENCE if s[0] == "scene"]


def lab_items():
    """The labs, in deck order, as (scene index they follow, file, title, (en, bn) notes)."""
    items, n_scene = [], -1
    for item in SEQUENCE:
        if item[0] == "scene":
            n_scene += 1
        else:
            name, title, notes = item[1:]
            items.append((n_scene, name, title, notes))
    return items


def convert_html(dest, folder, title, cdn):
    """One reveal.js page from the rendered slides, with the labs put after their chapters."""
    manim_slides(
        "convert", *scene_names(), str(dest),
        "--folder", str(folder),
        "--to", "html",
        "--use-template", str(HERE / "template" / "deck.html"),
        f"-ctitle={title}",
        "-cassets_dir=index_assets",
        "-ccontrols=true",
        "-cprogress=true",
        "-cslide_number=c/t",
        "-chash=true",
        "-cview_distance=2",
        # Bundle reveal.js into the deck, so it presents without an internet connection.
        *([] if cdn else ["--offline"]),
        check=True,
    )
    # Put each lab after the chapter it follows. The template leaves a marker after every
    # scene: <!-- after-scene:N --> with N counting scenes from 0.
    html = dest.read_text().replace("__AUTHOR__", AUTHOR)
    labs_after = {}
    for n, name, lab_title, notes in lab_items():
        labs_after.setdefault(n, []).append(lab_section(name, lab_title, notes))
    for n, sections in labs_after.items():
        marker = f"<!-- after-scene:{n} -->"
        if marker not in html:
            sys.exit(f"Template marker {marker} not found")
        html = html.replace(marker, marker + "\n" + "\n".join(sections))
    dest.write_text(html)


def build_deck(pptx=False, cdn=False):
    deck = BUILD / "deck"
    missing = [s for s in scene_names() if not (HERE / "slides" / f"{s}.json").exists()]
    if missing:
        sys.exit(f"Not rendered yet: {', '.join(missing)}. Run build.py without --deck-only first.")
    if deck.exists():
        shutil.rmtree(deck)
    deck.mkdir(parents=True)
    convert_html(deck / "index.html", HERE / "slides", TITLE, cdn)
    shutil.copytree(HERE / "labs", deck / "labs")
    shutil.copy(HERE / "template" / "speaker.html", deck / "speaker.html")
    thumbnails(deck)
    # The fonts, so the labs and the speaker view (whose notes are also in Bangla) look right
    # offline. Both look in labs/fonts/.
    (deck / "labs" / "fonts").mkdir()
    for f in [*FONTS.glob("Inter-*.ttf"), *FONTS.glob("NotoSansBengali-*.ttf"), *FONTS.glob("OFL*.txt")]:
        shutil.copy(f, deck / "labs" / "fonts" / f.name)
    print(f"Deck: {(deck / 'index.html').relative_to(HERE)}")
    if pptx:
        export_pptx(BUILD / "lecture3.pptx", HERE / "slides")


# --- PowerPoint ------------------------------------------------------------------------------
LAB_SHOTS = HERE / "images" / "labs"  # a screenshot of each lab, for PowerPoint


def export_pptx(out, folder):
    """PowerPoint, with the animations as videos and the notes in both languages. A lab can't
    run in PowerPoint, so its slide is a screenshot that says where the live one is."""
    from pptx import Presentation
    from pptx.util import Pt

    manim_slides("convert", *scene_names(), str(out), "--folder", str(folder), "--to", "pptx", check=True)
    prs = Presentation(str(out))
    counts = [len(json.loads((folder / f"{s}.json").read_text())["slides"]) for s in scene_names()]
    ids = prs.slides._sldIdLst
    # From the last lab back, so the earlier positions don't move.
    for n, name, title, (en, bn) in reversed(lab_items()):
        slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1])
        shot = LAB_SHOTS / (Path(name).stem + ".jpg")
        if shot.exists():
            slide.shapes.add_picture(str(shot), 0, 0, prs.slide_width, prs.slide_height)
        box = slide.shapes.add_textbox(Pt(18), prs.slide_height - Pt(40), prs.slide_width - Pt(36), Pt(30))
        box.text_frame.text = f"{title}: a screenshot. Open the HTML deck (build/{DECK}/index.html) to use it live."
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        slide.notes_slide.notes_text_frame.text = f"{en}\n\n{BANGLA}\n\n{bn}"
        ids.insert(sum(counts[: n + 1]), ids[-1])
    prs.save(str(out))
    print(f"PowerPoint: {out.relative_to(HERE)} ({len(prs.slides)} slides)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-q", "--quality", choices=QUALITY, default="l")
    ap.add_argument("--only", nargs="*", default=None, help="render only scenes whose class name contains one of these")
    ap.add_argument("--deck-only", action="store_true")
    ap.add_argument("--pptx", action="store_true")
    ap.add_argument("--cdn", action="store_true", help="load reveal.js from the internet instead of bundling it")
    ap.add_argument("-j", "--jobs", type=int, default=os.cpu_count() or 2)
    args = ap.parse_args()

    BUILD.mkdir(exist_ok=True)
    if not args.deck_only:
        todo = [(s[1], s[2]) for s in SEQUENCE if s[0] == "scene"]
        if args.only:
            todo = [t for t in todo if any(k.lower() in t[1].lower() for k in args.only)]
        print(f"Rendering {len(todo)} scene(s) at {QUALITY[args.quality][1]}, {args.jobs} at a time:")
        with ThreadPoolExecutor(args.jobs) as pool:
            ok = list(pool.map(lambda t: render(t[0], t[1], args.quality), todo))
        if not all(ok):
            sys.exit(1)
    build_deck(pptx=args.pptx, cdn=args.cdn)


if __name__ == "__main__":
    main()
