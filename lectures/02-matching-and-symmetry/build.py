#!/usr/bin/env python3
"""Render lecture 2 and build its slide deck: English slides, English and Bangla speaker notes.

    python build.py                  draft render (480p, 15 fps) of every chapter
    python build.py -q p             for presenting: 1080p, 30 fps
    python build.py --only M03       render only the matching chapters, then rebuild the deck
    python build.py --deck-only      rebuild the deck from what's already rendered
    python build.py --pptx           also export PowerPoint (the lab as a screenshot)
    python build.py --cdn            load reveal.js from the internet instead of bundling it

Quality: l 480p15 (draft), m 720p30, p 1080p30 (presenting), h 1080p60, k 4K60.

The deck lands in build/deck/. Open index.html in a browser: → and ← move between slides, S
opens the speaker view, F goes full screen. Run it from anywhere; paths are relative to this file.
"""

import argparse
import json
import os
import shutil
import subprocess
import sys
import time
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

HERE = Path(__file__).resolve().parent
FONTS = HERE.parents[1] / "kit" / "fonts"
BUILD = HERE / "build"
DECK = "deck"

TITLE = "Lecture 2 · Matching and Symmetry, in Layout"
AUTHOR = "Khandaker Ehsanul Karim"  # shown small on every slide: "Prepared by ..."
# Between a note's English and its Bangla, as kit.style.say() writes it.
BANGLA = "— বাংলা —"

PLACEMENT_LAB = (
    "lab",
    "lab1_placement.html",
    "Live lab: Place the units",
    (
        "Live lab. Each square is one unit device; A and B are the two halves of a matched pair. The "
        "colour behind them is a process gradient. Start with Side by side: A sits on one end of the "
        "gradient and B on the other, and the readout shows the mismatch. Press Interdigitated: smaller, "
        "not zero. Mirrored, A B B A: zero for a straight gradient, and turning the angle doesn't change "
        "that. Now add curvature: the long mirrored row leaves some error, and Cross-quad, the same "
        "units folded into two short rows, leaves almost none. Click any square to try your own "
        "arrangement. The bandgap's 1 : 8 array is the last preset.",
        "এইটা live lab। প্রত্যেকটা ঘর একটা unit device; A আর B একটা matched pair-এর দুই অর্ধেক। পেছনের রং একটা "
        "process gradient। Side by side দিয়া শুরু করেন: A gradient-এর এক মাথায়, B আরেক মাথায়, আর readout mismatch "
        "দেখায়। Interdigitated চাপেন: কম, কিন্তু শূন্য না। Mirrored, A B B A: সোজা gradient-এর জন্য শূন্য, আর angle "
        "ঘুরাইলেও সেইটা বদলায় না। এবার curvature যোগ করেন: লম্বা mirrored row কিছু error রাখে, আর Cross-quad, মানে একই "
        "unit দুইটা ছোট row-তে ভাঁজ করা, প্রায় কিছুই রাখে না। নিজের সাজানো চেষ্টা করতে যেকোনো ঘরে click করেন। শেষ "
        "preset-টা bandgap-এর 1 : 8 array।",
    ),
)

# The deck, in order: ("scene", file, class) for a Manim chapter, ("lab", file, title, (English
# notes, Bangla notes)) for a live lab.
SEQUENCE = [
    ("scene", "chapters/m00_cold_open.py", "M00ColdOpen"),
    ("scene", "chapters/m01_enemies.py", "M01Enemies"),
    ("scene", "chapters/m02_floorplan.py", "M02Floorplan"),
    ("scene", "chapters/m03_input_pair.py", "M03InputPair"),
    PLACEMENT_LAB,
    ("scene", "chapters/m04_mirrors.py", "M04Mirrors"),
    ("scene", "chapters/m05_cmfb.py", "M05CMFB"),
    ("scene", "chapters/m06_routing.py", "M06Routing"),
    ("scene", "chapters/m07_bandgap.py", "M07Bandgap"),
    ("scene", "chapters/m08_review.py", "M08Review"),
]

# Manim settings per quality. Use long options: manim-slides reads the "h" in "-qh" as its own
# -h (help) and exits without rendering.
QUALITY = {
    "l": (["--quality", "l"], "480p, 15 fps"),
    "m": (["--quality", "m"], "720p, 30 fps"),
    "p": (["--resolution", "1920,1080", "--frame_rate", "30"], "1080p, 30 fps"),
    "h": (["--quality", "h"], "1080p, 60 fps"),
    "k": (["--quality", "k"], "4K, 60 fps"),
}


def manim_slides(*args, **kw):
    return subprocess.run([sys.executable, "-m", "manim_slides", *args], cwd=HERE, **kw)


def render(scene_file, scene, quality):
    log = BUILD / f"{scene}.log"
    start = time.time()
    with open(log, "w") as out:
        proc = manim_slides(
            # Each scene gets its own media folder: renders run in parallel, and a shared text
            # cache lets one process read a file another is still writing.
            "render", "--media_dir", str(BUILD / "media" / scene), *QUALITY[quality][0], scene_file, scene,
            stdout=out, stderr=subprocess.STDOUT,
        )
    took = time.time() - start
    out_json = HERE / "slides" / f"{scene}.json"
    fresh = out_json.exists() and out_json.stat().st_mtime >= start
    ok = proc.returncode == 0 and fresh
    status = "ok" if ok else f"FAILED (see {log.relative_to(HERE)})"
    print(f"  {scene:<28} {took:6.0f} s  {status}", flush=True)
    return ok


def lab_section(name, title, notes):
    en, bn = notes
    return (
        '<section class="lab" data-background-color="#0E1116">'
        f'<iframe data-src="labs/{name}" title="{title}" allow="fullscreen"></iframe>'
        f'<aside class="notes">{en}\n\n{BANGLA}\n\n{bn}</aside>'
        "</section>"
    )


def thumbnails(deck):
    """The last frame of every slide's video, small, for the speaker view's "Next" preview."""
    out = deck / "thumbs"
    out.mkdir()
    for video in sorted((deck / "index_assets").glob("*.mp4")):
        jpg = out / (video.stem + ".jpg")
        for seek in (["-sseof", "-0.1"], []):  # a very short clip has no 0.1 s to seek back
            subprocess.run(["ffmpeg", "-v", "error", "-y", *seek, "-i", str(video), "-frames:v", "1",
                            "-vf", "scale=640:-2", "-q:v", "4", str(jpg)], check=False)
            if jpg.exists():
                break


def scene_names():
    return [s[2] for s in SEQUENCE if s[0] == "scene"]


def lab_items():
    """The labs, in deck order, as (scene index they follow, file, title, (en, bn) notes)."""
    items, n_scene = [], -1
    for item in SEQUENCE:
        if item[0] == "scene":
            n_scene += 1
        else:
            name, title, notes = item[1:]
            items.append((n_scene, name, title, notes))
    return items


def convert_html(dest, folder, title, cdn):
    """One reveal.js page from the rendered slides, with the labs put after their chapters."""
    manim_slides(
        "convert", *scene_names(), str(dest),
        "--folder", str(folder),
        "--to", "html",
        "--use-template", str(HERE / "template" / "deck.html"),
        f"-ctitle={title}",
        "-cassets_dir=index_assets",
        "-ccontrols=true",
        "-cprogress=true",
        "-cslide_number=c/t",
        "-chash=true",
        "-cview_distance=2",
        # Bundle reveal.js into the deck, so it presents without an internet connection.
        *([] if cdn else ["--offline"]),
        check=True,
    )
    # Put each lab after the chapter it follows. The template leaves a marker after every
    # scene: <!-- after-scene:N --> with N counting scenes from 0.
    html = dest.read_text().replace("__AUTHOR__", AUTHOR)
    labs_after = {}
    for n, name, lab_title, notes in lab_items():
        labs_after.setdefault(n, []).append(lab_section(name, lab_title, notes))
    for n, sections in labs_after.items():
        marker = f"<!-- after-scene:{n} -->"
        if marker not in html:
            sys.exit(f"Template marker {marker} not found")
        html = html.replace(marker, marker + "\n" + "\n".join(sections))
    dest.write_text(html)


def build_deck(pptx=False, cdn=False):
    deck = BUILD / "deck"
    missing = [s for s in scene_names() if not (HERE / "slides" / f"{s}.json").exists()]
    if missing:
        sys.exit(f"Not rendered yet: {', '.join(missing)}. Run build.py without --deck-only first.")
    if deck.exists():
        shutil.rmtree(deck)
    deck.mkdir(parents=True)
    convert_html(deck / "index.html", HERE / "slides", TITLE, cdn)
    shutil.copytree(HERE / "labs", deck / "labs")
    shutil.copy(HERE / "template" / "speaker.html", deck / "speaker.html")
    thumbnails(deck)
    # The fonts, so the labs and the speaker view (whose notes are also in Bangla) look right
    # offline. Both look in labs/fonts/.
    (deck / "labs" / "fonts").mkdir()
    for f in [*FONTS.glob("Inter-*.ttf"), *FONTS.glob("NotoSansBengali-*.ttf"), *FONTS.glob("OFL*.txt")]:
        shutil.copy(f, deck / "labs" / "fonts" / f.name)
    print(f"Deck: {(deck / 'index.html').relative_to(HERE)}")
    if pptx:
        export_pptx(BUILD / "lecture2.pptx", HERE / "slides")


# --- PowerPoint ------------------------------------------------------------------------------
LAB_SHOTS = HERE / "images" / "labs"  # a screenshot of each lab, for PowerPoint


def export_pptx(out, folder):
    """PowerPoint, with the animations as videos and the notes in both languages. A lab can't
    run in PowerPoint, so its slide is a screenshot that says where the live one is."""
    from pptx import Presentation
    from pptx.util import Pt

    manim_slides("convert", *scene_names(), str(out), "--folder", str(folder), "--to", "pptx", check=True)
    prs = Presentation(str(out))
    counts = [len(json.loads((folder / f"{s}.json").read_text())["slides"]) for s in scene_names()]
    ids = prs.slides._sldIdLst
    # From the last lab back, so the earlier positions don't move.
    for n, name, title, (en, bn) in reversed(lab_items()):
        slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[-1])
        shot = LAB_SHOTS / (Path(name).stem + ".jpg")
        if shot.exists():
            slide.shapes.add_picture(str(shot), 0, 0, prs.slide_width, prs.slide_height)
        box = slide.shapes.add_textbox(Pt(18), prs.slide_height - Pt(40), prs.slide_width - Pt(36), Pt(30))
        box.text_frame.text = f"{title}: a screenshot. Open the HTML deck (build/{DECK}/index.html) to use it live."
        box.text_frame.paragraphs[0].runs[0].font.size = Pt(14)
        slide.notes_slide.notes_text_frame.text = f"{en}\n\n{BANGLA}\n\n{bn}"
        ids.insert(sum(counts[: n + 1]), ids[-1])
    prs.save(str(out))
    print(f"PowerPoint: {out.relative_to(HERE)} ({len(prs.slides)} slides)")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("-q", "--quality", choices=QUALITY, default="l")
    ap.add_argument("--only", nargs="*", default=None, help="render only scenes whose class name contains one of these")
    ap.add_argument("--deck-only", action="store_true")
    ap.add_argument("--pptx", action="store_true")
    ap.add_argument("--cdn", action="store_true", help="load reveal.js from the internet instead of bundling it")
    ap.add_argument("-j", "--jobs", type=int, default=os.cpu_count() or 2)
    args = ap.parse_args()

    BUILD.mkdir(exist_ok=True)
    if not args.deck_only:
        todo = [(s[1], s[2]) for s in SEQUENCE if s[0] == "scene"]
        if args.only:
            todo = [t for t in todo if any(k.lower() in t[1].lower() for k in args.only)]
        print(f"Rendering {len(todo)} scene(s) at {QUALITY[args.quality][1]}, {args.jobs} at a time:")
        with ThreadPoolExecutor(args.jobs) as pool:
            ok = list(pool.map(lambda t: render(t[0], t[1], args.quality), todo))
        if not all(ok):
            sys.exit(1)
    build_deck(pptx=args.pptx, cdn=args.cdn)


if __name__ == "__main__":
    main()
